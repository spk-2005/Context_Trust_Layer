import csv
import io
import json
import logging
import re
import xml.etree.ElementTree as ET
import zipfile
from enum import Enum
from html.parser import HTMLParser
from typing import Any, Dict, List, Optional, Tuple
from pydantic import BaseModel, Field

logger = logging.getLogger(__name__)

_WORDSEGMENT_LOADED = False

def _get_wordsegment():
    global _WORDSEGMENT_LOADED
    try:
        import wordsegment
        if not _WORDSEGMENT_LOADED:
            wordsegment.load()
            _WORDSEGMENT_LOADED = True
        return wordsegment
    except Exception:
        return None

def repair_ocr_text(raw_text: str) -> str:
    """Repairs OCR-induced token concatenation, missing punctuation spacing, and run-together words."""
    if not raw_text:
        return ""

    ws = _get_wordsegment()

    lines = raw_text.splitlines()
    repaired_lines = []
    for line in lines:
        if not line.strip():
            repaired_lines.append(line)
            continue

        # Clean special quote / bracket artifacts
        line = line.replace("《", " ").replace("》", " ").replace("“", '"').replace("”", '"').replace("’", "'")
        line = re.sub(r"([a-zA-Z])'s([a-zA-Z])", r"\1's \2", line)

        # 1. Punctuation spacing: comma, period, parens, colon, semicolon
        line = re.sub(r'([a-zA-Z])\(', r'\1 (', line)
        line = re.sub(r'\)([a-zA-Z])', r') \1', line)
        line = re.sub(r',([a-zA-Z])', r', \1', line)
        line = re.sub(r':([a-zA-Z])', r': \1', line)
        line = re.sub(r';([a-zA-Z])', r'; \1', line)
        line = re.sub(r'\.([A-Za-z])', r'. \1', line)

        # 2. Letter to digit / digit to letter transitions (e.g. next30days -> next 30 days)
        line = re.sub(r'([a-zA-Z])(\d+)', r'\1 \2', line)
        line = re.sub(r'(\d+)([a-zA-Z])', r'\1 \2', line)

        tokens = line.split()
        repaired_tokens = []
        for token in tokens:
            if "@" in token or token.startswith(("http://", "https://", "www.")):
                repaired_tokens.append(token)
                continue

            match = re.match(r"^([^\w]*)([\w]+)([^\w]*)$", token)
            if not match:
                repaired_tokens.append(token)
                continue

            prefix, word, suffix = match.groups()
            if len(word) <= 3 or word.isdigit():
                repaired_tokens.append(token)
                continue

            # CamelCase splitting (e.g. SafeDrinkingWaterAct -> Safe Drinking Water Act)
            camel_parts = re.findall(r'[A-Z]+(?=[A-Z][a-z]|\d|\W|$)|[A-Z]?[a-z]+|\d+', word)
            if len(camel_parts) > 1 and all(len(p) >= 2 for p in camel_parts):
                repaired_tokens.append(prefix + " ".join(camel_parts) + suffix)
            elif len(word) >= 5 and ws is not None and not any(c.isdigit() for c in word):
                try:
                    segs = ws.segment(word.lower())
                    if len(segs) > 1 and all(len(s) >= 1 for s in segs):
                        if word.isupper():
                            segs = [s.upper() for s in segs]
                        elif word[0].isupper():
                            segs[0] = segs[0].capitalize()
                        repaired_tokens.append(prefix + " ".join(segs) + suffix)
                    else:
                        repaired_tokens.append(token)
                except Exception:
                    repaired_tokens.append(token)
            else:
                repaired_tokens.append(token)

        repaired_lines.append(" ".join(repaired_tokens))

    return "\n".join(repaired_lines)

_OCR_ENGINE = None

def get_ocr_engine():
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _OCR_ENGINE = RapidOCR()
            logger.info("RapidOCR engine initialized successfully")
        except Exception as e:
            logger.warning("RapidOCR could not be initialized: %s", e)
            _OCR_ENGINE = False
    return _OCR_ENGINE if _OCR_ENGINE is not False else None


class ContentBlockType(str, Enum):
    PARAGRAPH = "paragraph"
    HEADING = "heading"
    TABLE = "table"
    LIST_ITEM = "list_item"
    CAPTION = "caption"
    HEADER_FOOTER = "header_footer"
    IMAGE_OCR = "image_ocr"
    CODE_OR_CONFIG = "code_or_config"
    UNKNOWN = "unknown"


class ContentBlock(BaseModel):
    id: str = Field(default_factory=lambda: f"blk_{re.sub(r'[^a-zA-Z0-9]', '', str(io.BytesIO().getvalue() or ''))[:8]}")
    block_type: ContentBlockType = ContentBlockType.PARAGRAPH
    text: str
    raw_text: str = ""
    page_number: Optional[int] = None
    section_heading: Optional[str] = None
    table_data: Optional[Dict[str, Any]] = None  # {"headers": [...], "rows": [[...]], "row_dicts": [...]}
    start_char_offset: int = 0
    end_char_offset: int = 0
    metadata: Dict[str, Any] = Field(default_factory=dict)


class DocumentPage(BaseModel):
    page_number: int
    text: str
    char_count: int = 0
    is_scanned_image: bool = False
    blocks: List[ContentBlock] = Field(default_factory=list)


class ParsedDocument(BaseModel):
    filename: str
    title: str
    file_type: str
    parser_name: str = "UniversalFileParser"
    parser_version: str = "2.0.0"
    content_hash: str = ""
    full_text: str
    raw_text: str
    pages: List[DocumentPage] = Field(default_factory=list)
    blocks: List[ContentBlock] = Field(default_factory=list)
    tables: List[Dict[str, Any]] = Field(default_factory=list)
    extraction_quality: float = 1.0  # 0.0 to 1.0
    is_scanned: bool = False
    warnings: List[str] = Field(default_factory=list)
    metadata: Dict[str, Any] = Field(default_factory=dict)


class _HTMLToTextExtractor(HTMLParser):
    """Converts HTML documents to clean markdown-formatted text."""

    def __init__(self):
        super().__init__()
        self.result = []
        self.current_tag = None
        self.ignore_tags = {"script", "style", "head", "meta", "link", "noscript"}
        self.in_ignore = False
        self.table_row = []
        self.in_table = False

    def handle_starttag(self, tag, attrs):
        self.current_tag = tag.lower()
        if self.current_tag in self.ignore_tags:
            self.in_ignore = True
            return

        if self.current_tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            level = int(self.current_tag[1])
            self.result.append("\n\n" + "#" * level + " ")
        elif self.current_tag == "p":
            self.result.append("\n\n")
        elif self.current_tag == "li":
            self.result.append("\n- ")
        elif self.current_tag == "tr":
            self.in_table = True
            self.table_row = []
        elif self.current_tag == "br":
            self.result.append("\n")

    def handle_endtag(self, tag):
        t = tag.lower()
        if t in self.ignore_tags:
            self.in_ignore = False
        elif t == "tr" and self.table_row:
            self.result.append("\n| " + " | ".join(self.table_row) + " |")
            self.table_row = []
        elif t == "table":
            self.in_table = False
            self.result.append("\n\n")

    def handle_data(self, data):
        if self.in_ignore:
            return
        cleaned = data.strip()
        if not cleaned:
            return
        if self.in_table and self.current_tag in {"td", "th"}:
            self.table_row.append(cleaned)
        else:
            self.result.append(data)

    def get_text(self) -> str:
        raw = "".join(self.result)
        return re.sub(r"\n{3,}", "\n\n", raw).strip()


class UniversalFileParser:
    @classmethod
    def _get_ocr_engine(cls):
        return get_ocr_engine()

    """Universal parser capable of extracting structured text and canonical blocks from any document format."""

    @classmethod
    def parse_file(cls, filename: str, raw_bytes: bytes, doc_title: Optional[str] = None) -> str:
        """Parses any input file and returns structured Markdown text (Backward compatible)."""
        parsed = cls.parse_structured(filename, raw_bytes, doc_title)
        return parsed.full_text

    @classmethod
    def parse_structured(cls, filename: str, raw_bytes: bytes, doc_title: Optional[str] = None) -> ParsedDocument:
        """Parses any input file and returns a rich ParsedDocument with page, block, and table fidelity."""
        import hashlib
        name = filename or "uploaded_document"
        title = doc_title or name.rsplit(".", 1)[0].replace("-", " ").replace("_", " ").title()
        title = title.replace("\x00", "").strip()
        lower_name = name.lower()
        content_hash = hashlib.sha256(raw_bytes).hexdigest()

        # Content-Addressed Pipeline Cache Lookup (4-Factor Invariant Key)
        cache_key = None
        try:
            from app.caching.pipeline_cache import default_pipeline_cache
            config_hash = hashlib.sha256(b"ocr=true,table=true,headings=v2").hexdigest()[:16]
            cache_key = default_pipeline_cache.compute_pipeline_key(
                stage_name="universal_parser",
                source_sha256=content_hash,
                parser_version="v3.5",
                parser_config_hash=config_hash,
                normalizer_version="v1.2",
            )
            cached_doc = default_pipeline_cache.get_model("universal_parser", cache_key, ParsedDocument)
            if cached_doc is not None:
                return cached_doc
        except Exception:
            pass

        file_type = "plain_text"
        if lower_name.endswith(".pdf"):
            file_type = "pdf"
        elif lower_name.endswith((".docx", ".doc")):
            file_type = "docx"
        elif lower_name.endswith((".xlsx", ".xls")):
            file_type = "xlsx"
        elif lower_name.endswith((".pptx", ".ppt")):
            file_type = "pptx"
        elif lower_name.endswith(".csv"):
            file_type = "csv"
        elif lower_name.endswith(".tsv"):
            file_type = "tsv"
        elif lower_name.endswith((".json", ".jsonl")):
            file_type = "json"
        elif lower_name.endswith((".html", ".htm", ".xhtml")):
            file_type = "html"
        elif lower_name.endswith(".xml"):
            file_type = "xml"
        elif lower_name.endswith(".rtf"):
            file_type = "rtf"
        elif lower_name.endswith((".yaml", ".yml", ".toml", ".ini", ".conf", ".sql", ".log", ".env")):
            file_type = "code_or_config"

        pages: List[DocumentPage] = []
        blocks: List[ContentBlock] = []
        tables: List[Dict[str, Any]] = []
        warnings: List[str] = []
        extraction_quality = 1.0
        is_scanned = False
        full_text = ""

        try:
            if file_type == "pdf":
                full_text, pages, blocks, tables, warnings, is_scanned, extraction_quality = cls._parse_pdf_structured(
                    raw_bytes, title, filename
                )
            elif file_type == "docx":
                full_text, pages, blocks, tables, warnings = cls._parse_docx_structured(
                    raw_bytes, title, filename
                )
            elif file_type == "xlsx":
                full_text, pages, blocks, tables, warnings = cls._parse_xlsx_structured(
                    raw_bytes, title, filename
                )
            elif file_type in ("csv", "tsv"):
                delimiter = "\t" if file_type == "tsv" else None
                full_text = cls._parse_csv(raw_bytes, title, delimiter=delimiter)
                blocks.append(ContentBlock(
                    id="blk_csv_1",
                    block_type=ContentBlockType.TABLE,
                    text=full_text,
                    raw_text=full_text,
                ))
            elif file_type == "json":
                full_text = cls._parse_json(raw_bytes, title, is_jsonl=lower_name.endswith(".jsonl"))
                blocks.append(ContentBlock(
                    id="blk_json_1",
                    block_type=ContentBlockType.CODE_OR_CONFIG,
                    text=full_text,
                    raw_text=full_text,
                ))
            elif file_type in ("html", "xhtml"):
                full_text = cls._parse_html(raw_bytes, title)
                blocks.append(ContentBlock(
                    id="blk_html_1",
                    block_type=ContentBlockType.PARAGRAPH,
                    text=full_text,
                    raw_text=full_text,
                ))
            elif file_type == "xml":
                full_text = cls._parse_xml(raw_bytes, title)
                blocks.append(ContentBlock(
                    id="blk_xml_1",
                    block_type=ContentBlockType.CODE_OR_CONFIG,
                    text=full_text,
                    raw_text=full_text,
                ))
            elif file_type == "rtf":
                full_text = cls._parse_rtf(raw_bytes, title)
                blocks.append(ContentBlock(
                    id="blk_rtf_1",
                    block_type=ContentBlockType.PARAGRAPH,
                    text=full_text,
                    raw_text=full_text,
                ))
            elif file_type == "code_or_config":
                ext = lower_name.rsplit(".", 1)[-1]
                full_text = cls._parse_code_or_config(raw_bytes, title, ext)
                blocks.append(ContentBlock(
                    id="blk_code_1",
                    block_type=ContentBlockType.CODE_OR_CONFIG,
                    text=full_text,
                    raw_text=full_text,
                ))
            else:
                full_text = cls._parse_plain_text(raw_bytes, title, filename)
                blocks.append(ContentBlock(
                    id="blk_text_1",
                    block_type=ContentBlockType.PARAGRAPH,
                    text=full_text,
                    raw_text=full_text,
                ))

        except Exception as e:
            logger.warning("UniversalFileParser error for %s: %s", filename, e)
            warnings.append(f"Parsing error: {str(e)}")
            full_text = cls._parse_plain_text(raw_bytes, title, filename)
            blocks.append(ContentBlock(
                id="blk_fallback_1",
                block_type=ContentBlockType.PARAGRAPH,
                text=full_text,
                raw_text=full_text,
            ))

        full_text = full_text.replace("\x00", "").strip()
        if not full_text:
            full_text = f"# {title}\n\nIngested document '{filename}' ({len(raw_bytes)} bytes)."

        result_doc = ParsedDocument(
            filename=filename,
            title=title,
            file_type=file_type,
            content_hash=content_hash,
            full_text=full_text,
            raw_text=full_text,
            pages=pages,
            blocks=blocks,
            tables=tables,
            extraction_quality=extraction_quality,
            is_scanned=is_scanned,
            warnings=warnings,
            metadata={"file_size": len(raw_bytes), "block_count": len(blocks), "page_count": len(pages)},
        )
        if cache_key:
            try:
                from app.caching.pipeline_cache import default_pipeline_cache
                default_pipeline_cache.set("universal_parser", cache_key, result_doc, metadata={"filename": filename})
            except Exception:
                pass
        return result_doc

    @classmethod
    def _parse_pdf_structured(
        cls, raw_bytes: bytes, title: str, filename: str
    ) -> Tuple[str, List[DocumentPage], List[ContentBlock], List[Dict[str, Any]], List[str], bool, float]:
        pages: List[DocumentPage] = []
        blocks: List[ContentBlock] = []
        tables: List[Dict[str, Any]] = []
        warnings: List[str] = []
        is_scanned = False
        empty_pages = 0
        total_chars = 0

        fitz_doc = None
        try:
            try:
                import pymupdf
                fitz_doc = pymupdf.open(stream=raw_bytes, filetype="pdf")
            except Exception:
                fitz_doc = None

            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(raw_bytes))
            except Exception:
                reader = None

            pages_text = []
            from app.normalization.normalizer import DocumentNormalizer

            total_page_count = len(fitz_doc) if fitz_doc is not None else (len(reader.pages) if reader and reader.pages else 0)

            # Metadata check for Microsoft Print to PDF or virtual PDF printers that vectorize text
            doc_producer = ""
            if fitz_doc is not None and getattr(fitz_doc, "metadata", None):
                doc_producer = (fitz_doc.metadata.get("producer") or "").lower()
            is_print_to_pdf_doc = any(
                term in doc_producer for term in ("print to pdf", "microsoft", "pdfwrite", "skia", "cairo")
            )

            for page_num in range(1, total_page_count + 1):
                clean_page_text = ""

                # 1. Try PyMuPDF text extraction first (faster & handles standard/complex glyph layouts)
                drawings_count = 0
                if fitz_doc is not None and page_num <= len(fitz_doc):
                    try:
                        fitz_page = fitz_doc[page_num - 1]
                        clean_page_text = (fitz_page.get_text("text") or "").strip()
                        drawings_count = len(fitz_page.get_drawings())
                    except Exception:
                        pass

                # 2. Fallback to pypdf layout text extraction
                if not clean_page_text and reader is not None and page_num <= len(reader.pages):
                    try:
                        page = reader.pages[page_num - 1]
                        page_text = page.extract_text(extraction_mode="layout") or ""
                        clean_page_text = page_text.strip()
                    except Exception:
                        try:
                            clean_page_text = (reader.pages[page_num - 1].extract_text() or "").strip()
                        except Exception:
                            pass

                char_count = len(clean_page_text)
                total_chars += char_count

                # Scanned image or vector-drawn outline detection:
                # Trigger OCR if:
                # - Character count is very low (< 60 chars)
                # - Document was produced by Microsoft: Print To PDF or similar and char_count is low (< 300 chars)
                # - Page has many vector drawings (> 40 paths) and char_count is suspiciously low (< 250 chars)
                page_needs_ocr = (
                    char_count < 60
                    or (is_print_to_pdf_doc and char_count < 300)
                    or (drawings_count > 40 and char_count < 250)
                )
                ocr_recovered = False
                page_is_scanned = char_count < 40
                if page_needs_ocr:
                    ocr_engine = cls._get_ocr_engine()
                    # A. Try embedded image OCR if images exist in pypdf page
                    if ocr_engine is not None and reader is not None and page_num <= len(reader.pages):
                        try:
                            page = reader.pages[page_num - 1]
                            images = getattr(page, "images", [])
                            if images and len(images) > 0:
                                page_ocr_lines = []
                                for img in images:
                                    try:
                                        result, _ = ocr_engine(img.data)
                                        if result:
                                            lines = [line[1] for line in result if float(line[2]) > 0.4]
                                            page_ocr_lines.extend(lines)
                                    except Exception as img_err:
                                        logger.warning("OCR image extraction error on page %d: %s", page_num, img_err)

                                if page_ocr_lines:
                                    ocr_text = repair_ocr_text("\n".join(page_ocr_lines)).strip()
                                    if len(ocr_text) > char_count:
                                        clean_page_text = ocr_text
                                        total_chars = total_chars - char_count + len(clean_page_text)
                                        char_count = len(clean_page_text)
                                        ocr_recovered = True
                                        page_is_scanned = True
                                        is_scanned = True
                                        warnings.append(f"Page {page_num}: text recovered via embedded OCR ({char_count} chars).")
                        except Exception as ocr_err:
                            logger.warning("Embedded OCR processing error on page %d: %s", page_num, ocr_err)

                    # B. Canvas Rasterization Fallback (PyMuPDF): for vector-drawn PDFs (e.g. Microsoft Print to PDF), scanned documents, or outline text
                    if fitz_doc is not None and page_num <= len(fitz_doc) and ocr_engine is not None and (char_count < 250 or is_print_to_pdf_doc or drawings_count > 40):
                        try:
                            fitz_page = fitz_doc[page_num - 1]
                            pix = fitz_page.get_pixmap(dpi=96)
                            img_bytes = pix.tobytes("png")
                            result, _ = ocr_engine(img_bytes, use_cls=False)
                            if result:
                                lines = [line[1] for line in result if float(line[2]) > 0.4]
                                if lines:
                                    canvas_ocr_text = repair_ocr_text("\n".join(lines)).strip()
                                    if len(canvas_ocr_text) > char_count:
                                        clean_page_text = canvas_ocr_text
                                        total_chars = total_chars - char_count + len(clean_page_text)
                                        char_count = len(clean_page_text)
                                        ocr_recovered = True
                                        page_is_scanned = True
                                        is_scanned = True
                                        warnings.append(f"Page {page_num}: text recovered via canvas rasterization OCR ({char_count} chars).")
                                        logger.info(f"Page {page_num}/{total_page_count}: recovered {char_count} chars via canvas rasterization OCR")
                        except Exception as raster_err:
                            logger.warning("PyMuPDF canvas rasterization OCR error on page %d: %s", page_num, raster_err)

                if char_count < 40:
                    empty_pages += 1

                page_blocks: List[ContentBlock] = []
                if clean_page_text:
                    from app.normalization.document_restructurer import DocumentSemanticRestructurer

                    # 1. Convert spatial layout tables to Markdown pipe tables first
                    page_with_tables = DocumentSemanticRestructurer.extract_tables_from_layout(clean_page_text)

                    # 2. Reflow prose lines while strictly preserving Markdown table rows
                    reflowed_page = DocumentNormalizer.reflow_fragmented_lines(page_with_tables)
                    pages_text.append(f"## Page {page_num}\n{reflowed_page}")

                    # Extract paragraphs, tables, and headings into ContentBlocks
                    lines = reflowed_page.split("\n\n")
                    for b_idx, chunk in enumerate(lines):
                        chk = chunk.strip()
                        if not chk:
                            continue
                        is_table = bool(chk.startswith("|") and "\n|" in chk)
                        is_hdg = bool(re.match(r"^(#+|Section\s+\d+|Chapter\s+\d+|Article\s+\d+|Rule\s+\d+)", chk, re.IGNORECASE))
                        if not is_hdg and re.match(r"^\d+\.\d*\s+[A-Z]", chk):
                            first_line = chk.split("\n")[0].strip()
                            if len(first_line.split()) <= 8 and not first_line.endswith((".", ";")) and not any(m in first_line.lower() for m in [" shall ", " must ", " should ", " will ", " may "]):
                                is_hdg = True
                        b_type = ContentBlockType.TABLE if is_table else (ContentBlockType.HEADING if is_hdg else (ContentBlockType.IMAGE_OCR if ocr_recovered else ContentBlockType.PARAGRAPH))

                        table_meta = None
                        if is_table:
                            tbl_lines = [l for l in chk.split("\n") if l.strip().startswith("|")]
                            if len(tbl_lines) >= 2:
                                hdrs = [c.strip() for c in tbl_lines[0].split("|") if c.strip()]
                                data_r = [[c.strip() for c in l.split("|") if c.strip()] for l in tbl_lines[2:]]
                                table_meta = {"headers": hdrs, "rows": data_r, "page_number": page_num}
                                tables.append(table_meta)

                        blk = ContentBlock(
                            id=f"blk_p{page_num}_{b_idx+1}",
                            block_type=b_type,
                            text=chk,
                            raw_text=chk,
                            page_number=page_num,
                            table_data=table_meta,
                        )
                        page_blocks.append(blk)
                        blocks.append(blk)

                pages.append(DocumentPage(
                    page_number=page_num,
                    text=clean_page_text,
                    char_count=char_count,
                    is_scanned_image=page_is_scanned,
                    blocks=page_blocks,
                ))

            total_pages = max(1, total_page_count)
            if empty_pages == total_pages or (total_chars / total_pages) < 50:
                is_scanned = True
                warnings.append("Document appears to be a scanned image-based PDF. Text density is low.")
                extraction_quality = 0.4
            else:
                if is_scanned:
                    extraction_quality = max(0.75, min(0.95, 1.0 - (empty_pages / total_pages) * 0.4))
                else:
                    extraction_quality = max(0.6, min(1.0, 1.0 - (empty_pages / total_pages) * 0.5))

            if pages_text:
                full_text = f"# {title}\n\n" + "\n\n".join(pages_text)
                try:
                    from app.normalization.document_restructurer import DocumentSemanticRestructurer
                    full_text = DocumentSemanticRestructurer.restructure(full_text)
                except Exception as restruct_err:
                    logger.warning("DocumentSemanticRestructurer failed: %s", restruct_err)
            else:
                full_text = f"# {title}\n\nIngested PDF document {filename} ({total_page_count} pages)."

            if fitz_doc is not None:
                try:
                    fitz_doc.close()
                except Exception:
                    pass

            return full_text, pages, blocks, tables, warnings, is_scanned, extraction_quality

        except Exception as e:
            if fitz_doc is not None:
                try:
                    fitz_doc.close()
                except Exception:
                    pass
            warnings.append(f"PDF extraction warning: {str(e)}")
            plain = cls._parse_plain_text(raw_bytes, title, filename)
            return plain, [], [], [], warnings, False, 0.5

    @classmethod
    def _parse_docx_structured(
        cls, raw_bytes: bytes, title: str, filename: str
    ) -> Tuple[str, List[DocumentPage], List[ContentBlock], List[Dict[str, Any]], List[str]]:
        blocks: List[ContentBlock] = []
        tables: List[Dict[str, Any]] = []
        warnings: List[str] = []
        paragraphs_text = []

        try:
            import docx
            doc = docx.Document(io.BytesIO(raw_bytes))
            for p_idx, p in enumerate(doc.paragraphs, start=1):
                txt = p.text.strip()
                if not txt:
                    continue
                style_name = (p.style.name if hasattr(p, "style") and p.style else "").lower()
                is_heading = "heading" in style_name or bool(
                    re.match(r"^(Chapter\s+\d+|Page\s+\d+|Section\s+\d+|Control\s+rule|Policy\s+\d+)", txt, re.IGNORECASE)
                )
                b_type = ContentBlockType.HEADING if is_heading else ContentBlockType.PARAGRAPH
                blocks.append(ContentBlock(
                    id=f"blk_docx_{p_idx}",
                    block_type=b_type,
                    text=f"## {txt}" if is_heading else txt,
                    raw_text=txt,
                ))
                paragraphs_text.append(f"## {txt}" if is_heading else txt)

            for t_idx, table in enumerate(doc.tables, start=1):
                table_rows = []
                headers = []
                for r_idx, row in enumerate(table.rows):
                    cells = [cell.text.strip() for cell in row.cells]
                    if r_idx == 0:
                        headers = cells
                    table_rows.append(cells)

                if table_rows:
                    formatted_rows = ["| " + " | ".join(r) + " |" for r in table_rows]
                    t_text = "\n" + "\n".join(formatted_rows) + "\n"
                    paragraphs_text.append(t_text)
                    t_data = {"headers": headers, "rows": table_rows}
                    tables.append(t_data)
                    blocks.append(ContentBlock(
                        id=f"blk_table_{t_idx}",
                        block_type=ContentBlockType.TABLE,
                        text=t_text,
                        raw_text=t_text,
                        table_data=t_data,
                    ))

            full_text = f"# {title}\n\n" + "\n\n".join(paragraphs_text)
            return full_text, [], blocks, tables, warnings
        except Exception as e:
            warnings.append(f"DOCX parser error: {str(e)}")
            plain = cls._parse_plain_text(raw_bytes, title, filename)
            return plain, [], [], [], warnings

    @classmethod
    def _parse_xlsx_structured(
        cls, raw_bytes: bytes, title: str, filename: str
    ) -> Tuple[str, List[DocumentPage], List[ContentBlock], List[Dict[str, Any]], List[str]]:
        blocks: List[ContentBlock] = []
        tables: List[Dict[str, Any]] = []
        warnings: List[str] = []
        full_text = cls._parse_xlsx(raw_bytes, title, filename)
        blocks.append(ContentBlock(
            id="blk_xlsx_1",
            block_type=ContentBlockType.TABLE,
            text=full_text,
            raw_text=full_text,
        ))
        return full_text, [], blocks, tables, warnings

    @classmethod
    def _parse_pdf(cls, raw_bytes: bytes, title: str, filename: str) -> str:
        text, _, _, _, _, _, _ = cls._parse_pdf_structured(raw_bytes, title, filename)
        return text

    @classmethod
    def _parse_docx(cls, raw_bytes: bytes, title: str, filename: str) -> str:
        text, _, _, _, _ = cls._parse_docx_structured(raw_bytes, title, filename)
        return text

    @classmethod
    def _parse_xlsx(cls, raw_bytes: bytes, title: str, filename: str) -> str:
        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
                shared_strings = []
                if "xl/sharedStrings.xml" in zf.namelist():
                    ss_tree = ET.fromstring(zf.read("xl/sharedStrings.xml"))
                    ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    for si in ss_tree.findall(".//main:si", ns):
                        t_nodes = [t.text for t in si.findall(".//main:t", ns) if t.text]
                        shared_strings.append("".join(t_nodes))

                sheet_files = sorted([f for f in zf.namelist() if f.startswith("xl/worksheets/sheet") and f.endswith(".xml")])
                sheets_output = []

                for idx, sheet_file in enumerate(sheet_files, start=1):
                    sheet_tree = ET.fromstring(zf.read(sheet_file))
                    ns = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                    rows_data = []

                    for row in sheet_tree.findall(".//main:row", ns):
                        row_vals = []
                        for cell in row.findall(".//main:c", ns):
                            cell_type = cell.get("t")
                            v_node = cell.find("main:v", ns)
                            if v_node is not None and v_node.text:
                                val = v_node.text
                                if cell_type == "s" and val.isdigit():
                                    idx_val = int(val)
                                    val = shared_strings[idx_val] if idx_val < len(shared_strings) else val
                                row_vals.append(val.strip())
                            else:
                                t_node = cell.find("main:is/main:t", ns)
                                if t_node is not None and t_node.text:
                                    row_vals.append(t_node.text.strip())
                        if row_vals and any(row_vals):
                            rows_data.append("| " + " | ".join(row_vals) + " |")

                    if rows_data:
                        sheets_output.append(f"## Sheet {idx}\n\n" + "\n".join(rows_data))

                if sheets_output:
                    return f"# {title}\n\n" + "\n\n".join(sheets_output)
        except Exception:
            pass

        return cls._parse_plain_text(raw_bytes, title, filename)

    @classmethod
    def _parse_pptx(cls, raw_bytes: bytes, title: str, filename: str) -> str:
        try:
            with zipfile.ZipFile(io.BytesIO(raw_bytes)) as zf:
                slide_files = sorted(
                    [f for f in zf.namelist() if f.startswith("ppt/slides/slide") and f.endswith(".xml")],
                    key=lambda x: int(re.search(r"\d+", x).group()) if re.search(r"\d+", x) else 0,
                )
                slides_output = []
                ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

                for idx, sfile in enumerate(slide_files, start=1):
                    tree = ET.fromstring(zf.read(sfile))
                    texts = [t.text.strip() for t in tree.findall(".//a:t", ns) if t.text and t.text.strip()]
                    if texts:
                        slides_output.append(f"## Slide {idx}\n" + "\n".join(texts))

                if slides_output:
                    return f"# {title}\n\n" + "\n\n".join(slides_output)
        except Exception:
            pass

        return cls._parse_plain_text(raw_bytes, title, filename)

    @classmethod
    def _parse_csv(cls, raw_bytes: bytes, title: str, delimiter: Optional[str] = None) -> str:
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                decoded = raw_bytes.decode(enc)
                f = io.StringIO(decoded)
                reader = csv.reader(f, delimiter=delimiter) if delimiter else csv.reader(f)
                rows = []
                for idx, row in enumerate(reader):
                    clean_row = [c.strip() for c in row if c.strip()]
                    if clean_row:
                        rows.append("| " + " | ".join(clean_row) + " |")
                        if idx == 0:
                            rows.append("| " + " | ".join(["---"] * len(clean_row)) + " |")
                if rows:
                    return f"# {title}\n\n" + "\n".join(rows[:500])
            except UnicodeDecodeError:
                continue
        return cls._parse_plain_text(raw_bytes, title, "document.csv")

    @classmethod
    def _parse_json(cls, raw_bytes: bytes, title: str, is_jsonl: bool = False) -> str:
        for enc in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                decoded = raw_bytes.decode(enc).strip()
                if is_jsonl:
                    lines = [json.loads(line) for line in decoded.splitlines() if line.strip()]
                    formatted = json.dumps(lines[:100], indent=2)
                else:
                    data = json.loads(decoded)
                    formatted = json.dumps(data, indent=2)
                return f"# {title}\n\n```json\n{formatted}\n```"
            except Exception:
                continue
        return cls._parse_plain_text(raw_bytes, title, "data.json")

    @classmethod
    def _parse_html(cls, raw_bytes: bytes, title: str) -> str:
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                html_str = raw_bytes.decode(enc)
                extractor = _HTMLToTextExtractor()
                extractor.feed(html_str)
                extracted = extractor.get_text()
                if extracted:
                    return f"# {title}\n\n{extracted}"
            except Exception:
                continue
        return cls._parse_plain_text(raw_bytes, title, "document.html")

    @classmethod
    def _parse_xml(cls, raw_bytes: bytes, title: str) -> str:
        for enc in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                xml_str = raw_bytes.decode(enc)
                tree = ET.fromstring(xml_str)
                texts = [elem.text.strip() for elem in tree.iter() if elem.text and elem.text.strip()]
                if texts:
                    return f"# {title}\n\n" + "\n\n".join(texts[:300])
            except Exception:
                continue
        return cls._parse_plain_text(raw_bytes, title, "data.xml")

    @classmethod
    def _parse_rtf(cls, raw_bytes: bytes, title: str) -> str:
        text = raw_bytes.decode("latin-1", errors="ignore")
        clean_text = re.sub(r"\\[a-zA-Z0-9\-]+", " ", text)
        clean_text = re.sub(r"[{}\\]", " ", clean_text)
        clean_text = re.sub(r"\s+", " ", clean_text).strip()
        return f"# {title}\n\n{clean_text}"

    @classmethod
    def _parse_code_or_config(cls, raw_bytes: bytes, title: str, lang: str) -> str:
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                decoded = raw_bytes.decode(enc)
                return f"# {title}\n\n```{lang}\n{decoded}\n```"
            except UnicodeDecodeError:
                continue
        return cls._parse_plain_text(raw_bytes, title, f"code.{lang}")

    @classmethod
    def _parse_plain_text(cls, raw_bytes: bytes, title: str, filename: str) -> str:
        for enc in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                decoded = raw_bytes.decode(enc)
                return f"# {title}\n\n{decoded.strip()}"
            except UnicodeDecodeError:
                continue
        return f"# {title}\n\nBinary file {filename} ({len(raw_bytes)} bytes)."
